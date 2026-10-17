"""Restore official template from backup; re-fill row 7 only (no validation tampering)."""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import win32com.client as win32

BACKUP = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_backup.xlsm")
OUT = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm")
DATA_ROW = 7
IMAGE_COLS = list(range(21, 31))
LISTING_CREATE = "Create or Replace (Full Update)"


def read_row7(path: Path) -> dict[int, object]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Template"]
    cells = {c: ws.cell(DATA_ROW, c).value for c in range(1, ws.max_column + 1)
             if ws.cell(DATA_ROW, c).value is not None and str(ws.cell(DATA_ROW, c).value).strip()}
    wb.close()
    return cells


def main() -> None:
    saved = read_row7(OUT)
    for c in IMAGE_COLS:
        saved.pop(c, None)

    shutil.copy2(BACKUP, OUT)

    excel = win32.DispatchEx("Excel.Application")
    try:
        excel.DisplayAlerts = False
    except AttributeError:
        pass
    try:
        wb = excel.Workbooks.Open(str(OUT.resolve()))
        ws = wb.Worksheets("Template")

        # Order matters: B7 product type first → enables L browse dropdown
        ws.Cells(DATA_ROW, 2).Value = "PHONE_ACCESSORY"
        # C7: official list value only — do NOT touch Validation objects
        ws.Cells(DATA_ROW, 3).ClearContents()
        ws.Cells(DATA_ROW, 3).Value = LISTING_CREATE

        for c in IMAGE_COLS:
            ws.Cells(DATA_ROW, c).ClearContents()

        skip = {2, 3, *IMAGE_COLS}
        if 10 not in saved:
            saved[10] = "GTIN Exempt"
        for col, val in sorted(saved.items()):
            if col in skip:
                continue
            ws.Cells(DATA_ROW, col).Value = val

        wb.Save()
        wb.Close(SaveChanges=True)
    finally:
        excel.Quit()

    # verify header row 6 untouched
    wb2 = openpyxl.load_workbook(OUT, read_only=True, data_only=True)
    ws2 = wb2["Template"]
    print("r6 C official:", ws2.cell(6, 3).value)
    print("r7 C:", ws2.cell(7, 3).value)
    print("r7 B:", ws2.cell(7, 2).value)
    print("DL C4-6:", [wb2["Dropdown Lists"].cell(r, 3).value for r in range(4, 7)])
    wb2.close()
    print("Restored from backup + row7 fill. Size:", OUT.stat().st_size)


if __name__ == "__main__":
    main()
