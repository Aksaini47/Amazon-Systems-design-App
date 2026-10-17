"""Fill PHONE_ACCESSORY template: 25 listings + Mongo device codes + matched product images.

Uses Excel COM only (never openpyxl save on .xlsm).
"""
from __future__ import annotations

import csv
import re
import shutil
import time
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pythoncom
import win32com.client as win32
from bson import ObjectId
from pymongo import MongoClient

ROOT = Path(__file__).resolve().parent
REVIEW_CSV = ROOT / "bulk_create_25_review.csv"
CAT_REPORT = ROOT / "Category+Listings+Report_05-12-2026.xlsm"
BACKUP = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_backup.xlsm")
OUT = ROOT / "PHONE_ACCESSORY_bulk_25_filled.xlsm"
OUT_DOWNLOADS = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm")
MATCH_REPORT = ROOT / "bulk_25_image_match_report.csv"
DATA_START_ROW = 7

# New 237-col template map
COL_SKU = 1
COL_PT = 2
COL_ACTION = 3
COL_TITLE = 7
COL_BRAND = 9
COL_BROWSE = 12
COL_IMG_START = 21
COL_IMG_END = 30
COL_DESC = 31
COL_BULLET1 = 32
COL_KEYWORD = 37
COL_COMPAT = 45
COL_FULFILL = 101
COL_QTY = 102
COL_PRICE = 106
COL_MRP = 107

SCREEN_CATEGORIES = ("02 Display Assembly", "03 LCD Screen")
EMDASH = "\u2014"
IMAGE_COLS = list(range(COL_IMG_START, COL_IMG_END + 1))

BRAND_ALIASES = {"redmi": "xiaomi", "mi": "xiaomi", "iphone": "apple", "1+": "oneplus"}
BRAND_PREFIXES = [
    "samsung", "apple", "motorola", "oneplus", "oppo", "vivo", "realme", "xiaomi",
    "redmi", "poco", "mi", "tecno", "infinix", "honor", "nokia", "asus", "lava",
    "micromax", "lenovo", "sony", "lg", "huawei", "iphone", "google", "pixel", "nothing",
]

_VARIANT_SUFFIX_RE = re.compile(r"^(.+?)-(\d+)$")
_QUALIFIER_TOKENS = {"DS", "DSN", "DUAL", "SIM", "DDS", "TDS", "/DS", "BANDS", "GSM", "LTE", "5G", "4G", "3G"}

SCREEN_MATCH = {
    "OLED": {"OLED"},
    "Incell LCD": {"LCD", "Incell LCD", "TFT LCD", "TFT"},
    "LCD": {"LCD", "Incell LCD", "TFT LCD", "TFT"},
    "TFT": {"TFT", "LCD", "TFT LCD"},
    "Not Specified": {"LCD", "OLED", "TFT", "Incell LCD", "TFT LCD"},
}


@dataclass
class ListingAxes:
    screen: str
    grade: str
    has_frame: bool
    is_careog: bool


@dataclass
class ProductAxes:
    screen: str
    has_frame: bool
    is_careog: bool
    is_original: bool


@dataclass
class RowResult:
    sku: str
    stock_name: str
    title: str
    device_name: str = ""
    device_id: str = ""
    model_codes: str = ""
    product_name: str = ""
    image_count: int = 0
    flags: list[str] = field(default_factory=list)


def load_mongo_uri() -> str:
    text = Path(r"c:\Projects\Amazon Systems Design\.env").read_text(encoding="utf-8")
    m = re.search(r'MONGODB_URI="([^"]+)"', text)
    if not m:
        raise RuntimeError("MONGODB_URI not found in .env")
    return m.group(1)


def _norm(s: str) -> str:
    s = (s or "").lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _is_qualifier(c: str) -> bool:
    if not c:
        return True
    up = c.upper().strip("/").strip("-")
    if up in _QUALIFIER_TOKENS:
        return True
    return len(c) <= 3 and c.isalpha()


def _strip_variant(c: str) -> str:
    m = _VARIANT_SUFFIX_RE.match(c)
    if not m:
        return c
    return m.group(1) if re.search(r"\d", m.group(1)) else c


def _clean_code(c: str) -> str:
    if not c:
        return ""
    c = c.strip()
    if not c:
        return ""
    if c.upper() in ("CANCELLED", "TBA", "TBD", "N/A", "NONE", "UNKNOWN"):
        return ""
    if _is_qualifier(c):
        return ""
    return _strip_variant(c)


def _clean_codes(codes: list) -> list[str]:
    out, seen = [], set()
    for c in codes or []:
        s = _clean_code(str(c))
        if s and s not in seen:
            out.append(s)
            seen.add(s)
    return out


def load_mongo_lookup(db) -> dict:
    lookup: dict[str, list[dict]] = {}
    for d in db.device_model_lookup.find(
        {},
        {
            "brandSlug": 1,
            "deviceName": 1,
            "deviceNameLower": 1,
            "codes": 1,
            "finalAssignedCode": 1,
            "indiaModelCode": 1,
            "deviceId": 1,
        },
    ):
        slug = d.get("brandSlug") or ""
        if slug and d.get("deviceNameLower"):
            lookup.setdefault(slug, []).append(
                {
                    "dn": d.get("deviceName") or "",
                    "dn_lower": d.get("deviceNameLower") or "",
                    "codes": d.get("codes") or [],
                    "india": d.get("indiaModelCode") or "",
                    "final": d.get("finalAssignedCode") or "",
                    "device_id": str(d.get("deviceId") or ""),
                }
            )
    return lookup


def smart_lookup(mongo_lookup: dict, brand: str, model_phrase: str) -> dict | None:
    if not brand or not model_phrase:
        return None
    bs = BRAND_ALIASES.get(brand.lower(), brand.lower())
    devices = mongo_lookup.get(bs, [])
    if not devices:
        return None

    target = _norm(model_phrase)
    for p in BRAND_PREFIXES:
        if target.startswith(p + " "):
            target = target[len(p) + 1 :]
            break
    if bs == "apple" and not target.startswith("iphone"):
        target = "iphone " + target
    if brand.lower() in ("redmi", "poco", "mi") and not target.startswith(brand.lower()):
        target = brand.lower() + " " + target

    target_tokens = set(target.split())
    target_digits = {t for t in target.split() if any(c.isdigit() for c in t)}

    for d in devices:
        if d["dn_lower"] == target:
            india = d["india"] or d["final"] or ""
            codes = _clean_codes(d["codes"])
            if india and india not in codes:
                codes.insert(0, india)
            return {
                "dn": d["dn"],
                "device_id": d["device_id"],
                "all_codes": codes,
            }

    best = None
    for d in devices:
        dn_lower = d["dn_lower"]
        if len(d["dn"]) > 30 and re.search(r"\b[A-Z0-9]{4,}\b", d["dn"]):
            continue
        dn_norm = _norm(dn_lower)
        dn_tokens = set(dn_norm.split())
        dn_digits = {t for t in dn_norm.split() if any(c.isdigit() for c in t)}
        if target_digits and not (target_digits & dn_digits):
            continue
        target_only = target_digits - dn_digits
        dn_only = dn_digits - target_digits
        if target_only and dn_only:
            if not (target_digits <= dn_digits or dn_digits <= target_digits):
                continue
        digit_overlap = len(target_digits & dn_digits)
        token_overlap = len(target_tokens & dn_tokens)
        score = digit_overlap * 10000 + token_overlap * 100
        if dn_tokens == target_tokens:
            score += 50000
        if dn_lower.startswith(target + " "):
            score += 500
        score -= abs(len(dn_lower) - len(target)) * 2
        if best is None or score > best[0]:
            best = (score, d)

    if best:
        d = best[1]
        india = d["india"] or d["final"] or ""
        codes = _clean_codes(d["codes"])
        if india and india not in codes:
            codes.insert(0, india)
        return {"dn": d["dn"], "device_id": d["device_id"], "all_codes": codes}
    return None


TITLE_HEAD_RE = re.compile(
    r"compatible\s+for\s+(.+?)\s+"
    r"(?:\(|with\s+frame|without\s+frame|fingerprint|lcd|oled|amoled|in[\s-]?cell|careog|super\s+(?:retina|oled|amoled)|display)",
    re.IGNORECASE,
)


def extract_brand_model(title: str) -> tuple[str, str, str]:
    t = (title or "").strip()
    m = TITLE_HEAD_RE.search(t)
    head = m.group(1).strip() if m else t.replace("Compatible for ", "").strip()
    head = re.sub(r"\s*\([^)]*\)\s*$", "", head).strip()
    parts = head.split()
    brand = parts[0] if parts else ""
    primary = head.split("/")[0].strip() if "/" in head else head
    return brand, head, primary


def parse_listing_axes(title: str, bullets: list[str] | None = None) -> ListingAxes:
    t = (title or "").lower()
    btxt = " ".join(bullets or []).lower()
    combined = t + " " + btxt

    if "careog" in combined:
        grade = "CareOG"
    elif "with frame" in combined or "(with frame)" in combined:
        grade = "With Frame"
    else:
        grade = "Standard"

    if "oled" in combined or "amoled" in combined:
        screen = "OLED"
    elif "incell" in combined:
        screen = "Incell LCD"
    elif "tft" in combined:
        screen = "TFT"
    elif "lcd" in combined:
        screen = "LCD"
    else:
        screen = "Not Specified"

    has_frame = (
        grade == "With Frame"
        or "(with frame)" in combined
        or re.search(r"\bwith\s+frame\b", combined) is not None
    )
    return ListingAxes(screen=screen, grade=grade, has_frame=has_frame, is_careog=grade == "CareOG")


def parse_product_axes(name: str) -> ProductAxes:
    n = (name or "").lower()
    if "oled" in n or "amoled" in n or "super amoled" in n:
        screen = "OLED"
    elif "tft" in n:
        screen = "TFT"
    elif "incell" in n:
        screen = "Incell LCD"
    elif "lcd" in n:
        screen = "LCD"
    else:
        screen = "Not Specified"
    has_frame = bool(
        re.search(r"with\s+frame", n)
        or re.search(r"\bwith\s+frame\b", n)
        or re.search(r"\bwf\b", n)
        or re.search(r"assembly\s+with\s*$", n)
        or re.search(r"full\s+assembly\s+with\b", n)
    )
    is_careog = "careog" in n or "care og" in n
    is_original = "original" in n
    return ProductAxes(screen=screen, has_frame=has_frame, is_careog=is_careog, is_original=is_original)


def screen_products_for_device(db, device: dict, candidates: list[str]) -> list[dict]:
    oid = device["_id"]
    names = [device.get("modelName") or ""] + candidates
    regexes: list[str] = []
    for nm in names:
        tokens = [t for t in _norm(nm).split() if len(t) > 1][-4:]
        if tokens:
            regexes.append(".*".join(re.escape(t) for t in tokens))
    or_clauses: list[dict] = [
        {"primaryDeviceId": oid},
        {"compatibleDevices.deviceId": str(oid)},
    ]
    for pat in regexes[:4]:
        or_clauses.append({"name": {"$regex": pat, "$options": "i"}})
        or_clauses.append({"deviceModelName": {"$regex": pat, "$options": "i"}})

    q = {
        "category": {"$in": list(SCREEN_CATEGORIES)},
        "images.0": {"$exists": True},
        "$or": or_clauses,
    }
    seen: set[str] = set()
    out: list[dict] = []
    for p in db.products.find(q):
        pid = str(p.get("_id"))
        if pid not in seen:
            seen.add(pid)
            out.append(p)
    return out


def quality_match(listing: ListingAxes, product_name: str) -> bool:
    p = parse_product_axes(product_name)

    # Strict frame rule: Mongo WF but template not WF -> reject
    if p.has_frame and not listing.has_frame:
        return False
    if listing.has_frame and not p.has_frame:
        return False

    allowed = SCREEN_MATCH.get(listing.screen, {listing.screen})
    if p.screen not in allowed:
        return False

    if listing.is_careog:
        return p.is_careog or p.is_original

    if p.is_careog and not listing.is_careog:
        return False

    # OLED listing may match Original Super AMOLED premium panel
    if listing.screen == "OLED" and p.is_original and p.screen == "OLED":
        return True

    # Budget LCD listing should not pick premium Original
    if (
        listing.grade == "Standard"
        and listing.screen in ("LCD", "Incell LCD", "TFT")
        and p.is_original
        and not listing.is_careog
    ):
        return False

    return True


def fit_bp1(model: str, codes: list[str], max_chars: int = 255) -> str:
    if not codes:
        return (
            f"Compatible with {model} {EMDASH} Check Settings > About Phone for exact Model Number before ordering."
        )[:max_chars]
    template = (
        f"Compatible with {{model}} ({{codes}}) {EMDASH} "
        "Check Settings > About for exact Model Number before ordering."
    )
    full = template.format(model=model, codes=", ".join(codes))
    if len(full) <= max_chars:
        return full
    trimmed = list(codes)
    while trimmed:
        trimmed.pop()
        trial = template.format(model=model, codes=", ".join(trimmed))
        if len(trial) <= max_chars:
            return trial
    return (
        f"Compatible with {model} ({codes[0]}) {EMDASH} Check Settings > About for exact Model Number before ordering."
    )[:max_chars]


def clean_stock_model(stock_name: str) -> str:
    n = stock_name
    for w in [
        "Incell", "OLED", "AMOLED", "LCD", "CareOG", "Frame", "With Frame",
        "White", "Black", "Gold", "Super",
    ]:
        n = re.sub(rf"\b{w}\b", " ", n, flags=re.I)
    return re.sub(r"\s+", " ", n).strip()


def brand_prefix(brand: str, name: str) -> str:
    n = name
    if brand == "APPLE" or n.lower().startswith("iphone"):
        return "Apple iPhone " + re.sub(r"^iPhone\s+", "", n, flags=re.I)
    if brand == "SAMSUNG" or n.lower().startswith("galaxy"):
        return "Samsung Galaxy " + re.sub(r"^(?:Samsung\s+)?Galaxy\s+", "Galaxy ", n, flags=re.I)
    if brand == "REALME":
        return "Realme " + re.sub(r"^Realme\s+", "", n, flags=re.I)
    if brand == "OPPO":
        return "Oppo " + re.sub(r"^Oppo\s+", "", n, flags=re.I)
    if brand == "XIAOMI/REDMI/POCO":
        low = n.lower()
        if low.startswith("poco"):
            return "Poco " + re.sub(r"^Poco\s+", "", n, flags=re.I)
        if low.startswith("11x"):
            return "Xiaomi Mi 11X"
        return "Redmi " + re.sub(r"^Redmi\s+", "", n, flags=re.I)
    if brand == "MIX":
        low = n.lower()
        if "moto" in low:
            return "Motorola " + re.sub(r"^Moto\s+", "Moto ", n, flags=re.I)
        if "asus" in low:
            return "Asus " + re.sub(r"^Asus\s+", "", n, flags=re.I)
    return n


def compatible_candidates(brand: str, stock_name: str, title: str) -> list[str]:
    from_title = compatible_model_from_title(title)
    from_stock = brand_prefix(brand, clean_stock_model(stock_name))
    out: list[str] = []
    for c in (from_stock, from_title):
        if c and c not in out:
            out.append(c)
        if c and "next" in c.lower():
            alt = re.sub(r"next", "nxt", c, flags=re.I)
            if alt not in out:
                out.append(alt)
        if c and "rog phone" in c.lower():
            m = re.search(r"rog\s+phone\s*\d*", c, flags=re.I)
            if m and m.group(0) not in out:
                out.append(m.group(0))
    return out


def score_device_match(device: dict, compatible: str) -> int:
    dn = _norm(device.get("modelName") or "")
    cn = _norm(compatible)
    if not dn or not cn:
        return -1
    ct = set(cn.split())
    dt = set(dn.split())
    c_digits = {t for t in ct if any(ch.isdigit() for ch in t)}
    d_digits = {t for t in dt if any(ch.isdigit() for ch in t)}
    if c_digits and not (c_digits & d_digits):
        return -1
    score = len(ct & dt) * 10 + len(cn) - abs(len(dn) - len(cn))
    if cn == dn:
        score += 200
    elif cn in dn or dn in cn:
        score += 80
    return score


def find_device_by_compatible(db, candidates: list[str], brand_hint: str) -> dict | None:
    best_dev = None
    best_score = -1
    for compatible in candidates:
        if not compatible:
            continue
        q = re.escape(compatible.strip())
        exact = db.devices.find_one(
            {"modelName": {"$regex": f"^{q}$", "$options": "i"}, "isActive": {"$ne": False}}
        )
        if exact:
            sc = score_device_match(exact, compatible)
            if sc > best_score:
                best_score, best_dev = sc, exact

        tokens = [t for t in re.split(r"\s+", compatible.strip()) if len(t) > 1]
        if len(tokens) >= 2:
            pat = ".*".join(re.escape(t) for t in tokens[-4:])
            for doc in db.devices.find(
                {"modelName": {"$regex": pat, "$options": "i"}, "isActive": {"$ne": False}}
            ).limit(20):
                sc = score_device_match(doc, compatible)
                if sc > best_score:
                    best_score, best_dev = sc, doc

        if brand_hint and tokens:
            slug = BRAND_ALIASES.get(brand_hint.lower(), brand_hint.lower())
            last = tokens[-1]
            for doc in db.devices.find(
                {
                    "brandSlug": slug,
                    "modelName": {"$regex": re.escape(last), "$options": "i"},
                }
            ).limit(20):
                sc = score_device_match(doc, compatible)
                if sc > best_score:
                    best_score, best_dev = sc, doc

    return best_dev if best_score >= 10 else None


def device_name_ok(device: dict, compatible: str) -> bool:
    dn = _norm(device.get("modelName") or "")
    cn = _norm(compatible)
    if not dn or not cn:
        return False
    if dn == cn or cn in dn or dn in cn:
        return True
    ct = set(cn.split())
    dt = set(dn.split())
    overlap = ct & dt
    if len(overlap) < max(1, min(2, len(ct) - 1)):
        return False
    # Digit tokens in compatible must appear in device name (S21 FE vs S20 FE)
    c_digits = {t for t in ct if any(ch.isdigit() for ch in t)}
    d_digits = {t for t in dt if any(ch.isdigit() for ch in t)}
    if c_digits and not (c_digits & d_digits):
        return False
    return True


def resolve_device_and_codes(
    db,
    mongo_lookup: dict,
    brand: str,
    primary: str,
    candidates: list[str],
) -> tuple[dict | None, dict | None, list[str]]:
    lookup = smart_lookup(mongo_lookup, brand, primary)
    device = find_device_by_compatible(db, candidates, brand)

    lookup_device = None
    if lookup and lookup.get("device_id"):
        try:
            lookup_device = db.devices.find_one({"_id": ObjectId(lookup["device_id"])})
        except Exception:
            lookup_device = None

    if lookup_device and candidates and not device_name_ok(lookup_device, candidates[0]):
        lookup_device = None

    if device and lookup_device:
        if str(device["_id"]) != str(lookup_device["_id"]):
            if device_name_ok(lookup_device, candidates[0]) and not device_name_ok(device, candidates[0]):
                device = lookup_device
    elif not device and lookup_device:
        device = lookup_device

    codes: list[str] = []
    if lookup:
        codes = list(lookup.get("all_codes") or [])
    if device and not codes:
        codes = _clean_codes(device.get("modelCodes") or [])

    return device, lookup, codes


def pick_best_product(products: list[dict], listing: ListingAxes) -> dict | None:
    matched = [p for p in products if quality_match(listing, p.get("name", ""))]
    if not matched:
        return None

    def score(p: dict) -> tuple:
        pa = parse_product_axes(p.get("name", ""))
        screen_exact = 1 if pa.screen == listing.screen else 0
        img_n = len(p.get("images") or [])
        original_penalty = 0 if (listing.is_careog or pa.is_original == listing.is_careog) else (
            -1 if pa.is_original else 0
        )
        return (screen_exact, original_penalty, img_n)

    return max(matched, key=score)


def browse_mobile_displays(template: Path) -> str:
    wb = openpyxl.load_workbook(template, read_only=True, data_only=True)
    vv = wb["Valid Values"]
    for c in range(3, vv.max_column + 1):
        v = vv.cell(10, c).value
        if v and "Mobile Displays" in str(v) and "11180541031" in str(v):
            wb.close()
            return str(v)
    wb.close()
    raise ValueError("Mobile Displays browse node not found")


CAT_DONOR_ROW = 56


def load_donor_static() -> tuple[str | None, str | None]:
    """Single-row read from category report — avoid full-sheet index (hangs on large xlsm)."""
    if not CAT_REPORT.exists():
        return None, None
    wb = openpyxl.load_workbook(CAT_REPORT, read_only=True, data_only=True)
    ws = wb.active
    packer = ws.cell(CAT_DONOR_ROW, 121).value
    warranty = ws.cell(CAT_DONOR_ROW, 470).value
    wb.close()
    return (
        str(packer) if packer is not None else None,
        str(warranty) if warranty is not None else None,
    )


def compatible_model_from_title(title: str) -> str:
    brand, head, _ = extract_brand_model(title)
    # Strip quality parens already done in head extraction
    return head.strip()


def build_row_cells(
    row: dict,
    browse: str,
    packer: str | None,
    warranty: str | None,
    mongo_lookup: dict,
    db,
) -> tuple[dict[int, object], RowResult]:
    sku = row["new_sku"]
    title = row["title"]
    stock_name = row.get("stock_name", "")
    sell = float(row["sell_price"]) if row.get("sell_price") else None
    mrp = float(row["mrp"]) if row.get("mrp") else None
    qty = int(float(row.get("qty") or 100))

    brand, _, primary = extract_brand_model(title)
    candidates = compatible_candidates(brand, stock_name, title)
    compatible = candidates[0] if candidates else compatible_model_from_title(title)

    device, lookup, codes = resolve_device_and_codes(db, mongo_lookup, brand, primary, candidates)

    res = RowResult(sku=sku, stock_name=stock_name, title=title)
    if lookup:
        res.device_name = (device or {}).get("modelName") or lookup.get("dn") or ""
    elif device:
        res.device_name = device.get("modelName") or ""
    if device:
        res.device_id = str(device.get("_id") or "")
    if lookup and lookup.get("device_id") and not res.device_id:
        res.device_id = lookup["device_id"]

    if not codes:
        res.flags.append("no_model_codes")

    model = compatible
    bp1 = fit_bp1(model, codes)
    res.model_codes = ", ".join(codes[:12])

    # Bullets from title tier (keep 2-5 from title signals)
    axes = parse_listing_axes(title)
    tier = "gold" if axes.is_careog else "silver" if axes.screen == "OLED" else "bronze"
    fp = "fingerprint support" in title.lower() and "no fingerprint" not in title.lower()
    tier_txt = {
        "gold": "Gold Tier CareOG — Tested OG combo with premium quality for professional repair.",
        "silver": "Silver Tier OLED — Premium display with vivid colors and better contrast.",
        "bronze": "Bronze Tier LCD — Budget friendly good quality display. Reliable performance.",
    }
    frame_note = " With frame for easy installation." if axes.has_frame else ""
    bullets = [
        bp1,
        tier_txt[tier],
        f"Quality Tested — Checked for dead pixels, color calibration, touch response.{frame_note}",
        "Professional installation recommended — Self-install without skills may damage device.",
        "7-Day replacement warranty — QC tested before dispatch. Film removal = acceptance.",
    ]

    images: list[str] = []
    if device:
        products = screen_products_for_device(db, device, candidates)
        best = pick_best_product(products, axes)
        if best:
            res.product_name = best.get("name", "")[:120]
            images = list(best.get("images") or [])[:9]
            res.image_count = len(images)
        else:
            res.flags.append("no_quality_matched_product")
    else:
        res.flags.append("no_device_match")

    kw = re.sub(r"[^a-z0-9 ]", " ", title.lower())
    kw = re.sub(r"\s+", " ", kw).strip()[:200]

    cells: dict[int, object] = {
        COL_SKU: sku,
        COL_PT: "PHONE_ACCESSORY",
        COL_ACTION: "Create or Replace (Full Update)",
        COL_TITLE: title[:200],
        COL_BRAND: "Generic",
        10: "GTIN Exempt",
        COL_BROWSE: browse,
        17: sku,
        18: "Generic",
        COL_DESC: (
            "Compatible mobile phone display and touch screen combo replacement. "
            "Tested for quality before dispatch. Professional installation recommended."
        ),
        COL_KEYWORD: kw,
        39: 1,
        40: "Mobile Screen Replacement Part",
        43: sku,
        COL_COMPAT: compatible,
        49: 1,
        50: "Piece",
        52: "Mobile Screen",
        57: "HSN Code",
        58: "851770",
        63: packer,
        68: packer,
        74: 25,
        75: "Grams",
        77: "New",
        79: "A_GEN_STANDARD",
        81: 1,
        COL_FULFILL: "AMAZON_IN",
        COL_QTY: qty,
        103: 2,
        105: "Disabled",
        134: 26,
        135: "Centimetres",
        136: 12,
        137: "Centimetres",
        138: 2,
        139: "Centimetres",
        140: 26,
        141: "Centimetres",
        142: 12,
        143: "Centimetres",
        144: 2,
        145: "Centimetres",
        146: 220,
        147: "Grams",
        149: "China",
        150: warranty,
        153: "No",
        176: "Not Applicable",
    }
    if sell:
        cells[COL_PRICE] = sell
    if mrp:
        cells[COL_MRP] = mrp
    for i, b in enumerate(bullets):
        cells[COL_BULLET1 + i] = str(b)[:255]
    for i, url in enumerate(images[:10]):
        cells[COL_IMG_START + i] = url

    return cells, res


def write_template(rows: list[tuple[dict[int, object], RowResult]]) -> Path:
    """Excel COM write — same pattern as prepare_phone_accessory_upload.py."""
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(BACKUP, OUT)

    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    try:
        try:
            excel.DisplayAlerts = False
        except AttributeError:
            pass
        wb = excel.Workbooks.Open(str(OUT.resolve()))
        time.sleep(1.5)
        ws = wb.Worksheets("Template")
        for offset, (cells, _) in enumerate(rows):
            r = DATA_START_ROW + offset
            # B before C/L — dropdown dependency (per FLOW.md)
            ws.Cells(r, COL_PT).Value = cells[COL_PT]
            ws.Cells(r, COL_ACTION).Value = cells[COL_ACTION]
            ws.Cells(r, COL_BROWSE).Value = cells[COL_BROWSE]
            skip = {COL_PT, COL_ACTION, COL_BROWSE}
            for col in IMAGE_COLS:
                ws.Cells(r, col).ClearContents()
            for col, val in sorted(cells.items()):
                if col in skip or val is None or str(val).strip() == "":
                    continue
                ws.Cells(r, col).Value = val
            time.sleep(0.15)
        wb.Save()
        wb.Close(SaveChanges=True)
    finally:
        try:
            excel.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()

    # Best-effort copy to Downloads if not locked
    try:
        if OUT_DOWNLOADS.exists():
            OUT_DOWNLOADS.unlink()
        shutil.copy2(OUT, OUT_DOWNLOADS)
        return OUT_DOWNLOADS
    except OSError:
        return OUT


def main() -> None:
    if not REVIEW_CSV.exists():
        raise FileNotFoundError(REVIEW_CSV)
    if not BACKUP.exists():
        raise FileNotFoundError(BACKUP)

    listings = list(csv.DictReader(REVIEW_CSV.open(encoding="utf-8")))
    print(f"Listings to fill: {len(listings)}")

    uri = load_mongo_uri()
    db = MongoClient(uri, serverSelectionTimeoutMS=25000)["repairfully"]
    mongo_lookup = load_mongo_lookup(db)
    print(f"Mongo devices loaded: {sum(len(v) for v in mongo_lookup.values())}")

    browse = browse_mobile_displays(BACKUP)
    print("Browse node OK", flush=True)
    packer, warranty = load_donor_static()
    print("Donor static OK", flush=True)

    built: list[tuple[dict[int, object], RowResult]] = []
    report: list[dict] = []
    for row in listings:
        cells, res = build_row_cells(row, browse, packer, warranty, mongo_lookup, db)
        built.append((cells, res))
        report.append(
            {
                "sku": res.sku,
                "stock_name": res.stock_name,
                "device": res.device_name,
                "model_codes": res.model_codes,
                "matched_product": res.product_name,
                "images": res.image_count,
                "flags": " | ".join(res.flags),
            }
        )
        status = "OK" if res.image_count else "MISSING"
        print(f"  [{status}] {res.sku}: device={res.device_name[:30] if res.device_name else '-'} imgs={res.image_count} {res.flags}")

    saved_path = write_template(built)

    with MATCH_REPORT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(report[0].keys()))
        w.writeheader()
        w.writerows(report)

    matched = sum(1 for r in report if int(r["images"]) > 0)
    print(f"\nSaved template: {saved_path}")
    print(f"Rows: {DATA_START_ROW}-{DATA_START_ROW + len(built) - 1}")
    print(f"Images matched: {matched}/{len(report)}")
    print(f"Report: {MATCH_REPORT}")


if __name__ == "__main__":
    main()
