"""Generate bulk CREATE file for 25 missing-quality stock SKUs."""
from __future__ import annotations

import csv
import re
import statistics
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
MISSING_CSV = ROOT / "stock_missing_this_quality.csv"
CAT_REPORT = ROOT / "Category+Listings+Report_05-12-2026.xlsm"
ALL_LISTINGS = ROOT / "All+Listings+Report_05-12-2026.txt"
STOCK_CSVS = [
    ROOT / "stock_1_exact_match.csv",
    ROOT / "stock_2_loose_match.csv",
    ROOT / "stock_3_not_listed.csv",
    ROOT / "stock_A_combined.csv",
    ROOT / "stock_wrong_quality.csv",
]
OUTPUT_XLSM = ROOT / "bulk_create_25_quality.xlsm"
REVIEW_CSV = ROOT / "bulk_create_25_review.csv"

# Category Listings Report column indices (row 4 headers)
COL_STATUS = 1
COL_TITLE = 2
COL_SKU = 3
COL_PRODUCT_TYPE = 4
COL_LISTING_ACTION = 5
COL_ITEM_NAME = 9
COL_BRAND = 10
COL_MAIN_IMAGE = 29
COL_DESCRIPTION = 39
COL_BULLET1 = 40
COL_BULLET5 = 44
COL_KEYWORD1 = 45
COL_QTY = 420
COL_PRICE = 424
COL_MRP = 425

GOOD_SUNSKY_TYPES = {
    "LCD Screen", "For Samsung", "For Xiaomi", "For OPPO", "For Nokia",
    "For OnePlus", "For Huawei", "For Motorola", "For iPhone",
}

# Manual cost overrides when stock CSV Sunsky match is wrong part type (INR)
SUNSKY_COST_OVERRIDES: dict[str, float] = {
    "Realme C53 Incell": 943,
    "Oppo F5 Frame": 943,
    "Galaxy J7 Next Incell": 265,
    "Galaxy J7 2015 Incell": 265,
    "Redmi 12 With Frame": 825,
    "Moto G34 Frame": 825,
    "Asus ROG Phone 2 Incell": 1213,
    "Oppo F17 Pro OLED": 1298,
}

SKIP_SUNSKY_TYPES = {
    "Back Cover", "Camera Series", "Flex Cable", "Outer Glass Lens",
    "Card Tray", "Touch Panel", "Frame Bezel Plate", "Charging Port Board",
    "Others", "Small Board", "Speaker Ringer Buzzer",
}

DEFAULT_MARKUP = {"careog": 4.5, "oled": 5.5, "lcd": 4.52}


def parse_rs(value: str | None) -> float | None:
    if not value:
        return None
    s = str(value).strip()
    s = re.sub(r"^Rs\.?", "", s, flags=re.I).replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        m = re.search(r"\d+\.?\d*", s)
        return float(m.group()) if m else None


def calculate_mrp(price: float) -> float:
    if price < 500:
        return round(price * 5, -2)
    if price <= 2000:
        return round(price * 2, -2)
    return round(price * 1.43, -2)


def load_sunsky_index() -> dict[str, dict]:
    """Index by stock item name; prefer display-related Sunsky product types."""
    idx: dict[str, dict] = {}

    def score(row: dict) -> int:
        pt = row.get("Sunsky Product Type", "")
        if pt in GOOD_SUNSKY_TYPES:
            return 3
        if pt in SKIP_SUNSKY_TYPES:
            return 0
        return 1

    for path in STOCK_CSVS:
        if not path.exists():
            continue
        with path.open(newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                item = row.get("Stock Item", "").strip()
                if not item or not row.get("Sunsky Price (INR)", "").strip():
                    continue
                if item not in idx or score(row) > score(idx[item]):
                    idx[item] = row
    return idx


def load_existing_skus() -> set[str]:
    skus: set[str] = set()
    if not ALL_LISTINGS.exists():
        return skus
    for line in ALL_LISTINGS.read_text(encoding="utf-8-sig").splitlines()[1:]:
        cols = line.split("\t")
        if len(cols) > 3 and cols[3].strip():
            skus.add(cols[3].strip().upper())
    return skus


def tier_key(screen: str, quality: str) -> str:
    if quality == "CareOG":
        return "careog"
    if screen == "OLED":
        return "oled"
    return "lcd"


def resolve_sell_price(
    name: str,
    screen: str,
    quality: str,
    sunsky_row: dict,
    markups: dict[str, float],
) -> tuple[float, float, list[str]]:
    flags: list[str] = []
    tk = tier_key(screen, quality)
    markup = markups.get(tk, DEFAULT_MARKUP[tk])

    if name in SUNSKY_COST_OVERRIDES:
        cost = SUNSKY_COST_OVERRIDES[name]
        flags.append("manual_sunsky_override")
        return round(cost * markup, 0), cost, flags

    sunsky_type = sunsky_row.get("Sunsky Product Type", "")
    cost = parse_rs(sunsky_row.get("Sunsky Price (INR)"))
    amazon_price = parse_rs(sunsky_row.get("Amazon Price"))

    if cost and sunsky_type in GOOD_SUNSKY_TYPES:
        return round(cost * markup, 0), cost, flags

    if amazon_price:
        mult = 1.0
        donor_q = sunsky_row.get("Listing Quality") or sunsky_row.get("Stock Quality", "")
        donor_s = sunsky_row.get("Listing Screen") or sunsky_row.get("Stock Screen", "")
        if quality == "Standard" and donor_q == "CareOG":
            mult = 0.78
        elif screen == "OLED" and "LCD" in str(donor_s):
            mult = 1.12
        elif quality == "With Frame" and donor_q != "With Frame":
            mult = 1.08
        elif screen == "Incell LCD" and donor_q == "CareOG":
            mult = 0.82
        flags.append("price_from_amazon_donor")
        return round(amazon_price * mult, 0), cost or 0, flags

    if cost:
        flags.append(f"suspect_sunsky_type:{sunsky_type}")
        return round(cost * markup, 0), cost, flags

    flags.append("missing_price")
    return 0, 0, flags


def compute_markups() -> dict[str, float]:
    buckets: dict[str, list[float]] = {"careog": [], "oled": [], "lcd": []}
    path = ROOT / "stock_1_exact_match.csv"
    if not path.exists():
        return DEFAULT_MARKUP.copy()
    with path.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            pt = row.get("Sunsky Product Type", "")
            if pt in SKIP_SUNSKY_TYPES:
                continue
            cost = parse_rs(row.get("Sunsky Price (INR)"))
            price = parse_rs(row.get("Amazon Price"))
            if not cost or not price or cost < 150:
                continue
            ratio = price / cost
            if ratio < 1.8 or ratio > 10:
                continue
            buckets[tier_key(row.get("Stock Screen", ""), row.get("Stock Quality", ""))].append(ratio)
    out = DEFAULT_MARKUP.copy()
    for k, vals in buckets.items():
        if vals:
            out[k] = round(statistics.median(vals), 2)
    return out


def brand_prefix(brand: str, name: str) -> str:
    n = name
    if brand == "APPLE" or n.lower().startswith("iphone"):
        return "Apple iPhone " + re.sub(r"^iPhone\s+", "", n, flags=re.I)
    if brand == "SAMSUNG" or n.lower().startswith("galaxy"):
        return "Samsung " + re.sub(r"^Galaxy\s+", "Galaxy ", n, flags=re.I)
    if brand == "ONEPLUS":
        return "OnePlus " + re.sub(r"^OnePlus\s+", "", n, flags=re.I)
    if brand == "REALME":
        return "Realme " + re.sub(r"^Realme\s+", "", n, flags=re.I)
    if brand == "OPPO":
        return "Oppo " + re.sub(r"^Oppo\s+", "", n, flags=re.I)
    if brand == "VIVO":
        return "Vivo " + re.sub(r"^Vivo\s+", "", n, flags=re.I)
    if brand == "XIAOMI/REDMI/POCO":
        if n.lower().startswith(("redmi", "poco", "11x", "a3")):
            if n.lower().startswith("11x"):
                return "Xiaomi Mi 11X"
            if n.lower().startswith("poco"):
                return "Poco " + re.sub(r"^Poco\s+", "", n, flags=re.I)
            return "Redmi " + re.sub(r"^Redmi\s+", "", n, flags=re.I)
        return "Xiaomi " + n
    if brand == "MIX":
        low = n.lower()
        if "moto" in low:
            return "Motorola " + re.sub(r"^Moto\s+", "Moto ", n, flags=re.I)
        if "honor" in low:
            return "Honor " + re.sub(r"^Honor\s+", "", n, flags=re.I)
        if "nokia" in low:
            return "Nokia " + re.sub(r"^Nokia\s+", "", n, flags=re.I)
        if "asus" in low:
            return "Asus " + re.sub(r"^Asus\s+", "", n, flags=re.I)
        if "nothing" in low:
            return "Nothing Phone " + re.sub(r"^Nothing Phone\s+", "", n, flags=re.I)
        if "infinix" in low:
            return "Infinix " + re.sub(r"^Infinix\s+", "", n, flags=re.I)
    return n


def clean_model_name(name: str) -> str:
    n = name
    for w in [
        "Incell", "OLED", "AMOLED", "LCD", "CareOG", "Frame", "With Frame",
        "White", "Black", "Gold", "Super",
    ]:
        n = re.sub(rf"\b{w}\b", " ", n, flags=re.I)
    n = re.sub(r"\s+5G\s+", " 5G ", n)
    n = re.sub(r"\s+4G\s+", " 4G ", n)
    return re.sub(r"\s+", " ", n).strip()


def build_title(brand: str, name: str, screen: str, quality: str) -> str:
    device = brand_prefix(brand, clean_model_name(name))
    fp = screen == "OLED" and quality != "CareOG"
    parts = [f"Compatible for {device}"]

    if quality == "With Frame":
        parts.append("(with Frame)")

    if quality == "CareOG":
        if fp:
            parts.append("(Fingerprint Support) CareOG")
        else:
            parts.append("CareOG")
    elif screen == "OLED":
        parts.append("(Fingerprint Support) OLED" if fp else "OLED")
    elif screen in ("Incell LCD", "LCD"):
        parts.append("(No Fingerprint Support) LCD" if not fp else "LCD")
    else:
        parts.append("LCD")

    parts.append("Display+Touch Screen Combo Folder")
    title = " ".join(parts)
    return title[:200]


def detect_fingerprint(screen: str, quality: str, title: str) -> bool:
    if "no fingerprint" in title.lower():
        return False
    if quality == "CareOG" or screen == "OLED":
        return True
    return False


def make_bullets(title: str, screen: str, quality: str, model: str) -> list[str]:
    frame = quality == "With Frame" or "with frame" in title.lower() or "folder" in title.lower()
    fp = detect_fingerprint(screen, quality, title)
    tier = tier_key(screen, quality)

    fp_txt = " — Fingerprint supported after fitting" if fp else ""
    b1 = f"Compatible with {model}{fp_txt} — Verify model number from Settings before ordering"

    if tier == "careog":
        b2 = "CareOG tested combo — Premium quality display with original chip IC. Reliable for professional repair technicians."
    elif tier == "oled":
        b2 = "OLED display — Premium panel with vivid colors and better contrast." + (" Fingerprint supported after fitting." if fp else "")
    else:
        b2 = "LCD display — Budget friendly good quality display. Great value for money, reliable performance."

    frame_note = " With frame for easy installation." if frame else ""
    if tier == "careog":
        b3 = f"Quality tested — 100% checked for dead pixels, color calibration, touch response.{frame_note}"
    elif tier == "oled":
        b3 = f"OLED quality — Original chip IC, vibrant colors, tested panel.{frame_note}"
    else:
        b3 = f"LCD quality — Good color reproduction, original chip compatible, tested.{frame_note}"

    if frame:
        b4 = (
            "With frame assembly — Screen pre-pasted on frame, ready to install. Original adhesive pre-applied."
            if not fp
            else "With frame assembly — Pre-framed screen for fingerprint calibration. Ready to install."
        )
    else:
        b4 = "Screen only — Requires frame transfer from old screen. Adhesive sheet included. Professional installation recommended."

    b5 = "Warranty: 7 days replacement — QC tested before dispatch. No warranty after film removal. For help visit Repairfully.com"
    return [b1, b2, b3, b4, b5]


def make_keywords(screen: str, quality: str, model: str) -> str:
    base = "display sceen replacement mobile repair marammat"
    if quality == "CareOG":
        tier_kw = "careog tested combo premium quality"
    elif screen == "OLED":
        tier_kw = "oled amoled premium fingerprint"
    else:
        tier_kw = "lcd budget friendly non fingerprint"
    kw = f"{base} {tier_kw} {model.lower()}"
    return kw[:200]


def make_sku(brand: str, name: str, screen: str, quality: str, existing: set[str]) -> str:
    clean = re.sub(r"[^A-Za-z0-9]", "", clean_model_name(name).upper())
    prefix_map = {
        "APPLE": "IP",
        "SAMSUNG": "SAM",
        "REALME": "RM",
        "OPPO": "OP",
        "VIVO": "VV",
        "ONEPLUS": "OP",
        "XIAOMI/REDMI/POCO": "RD",
        "MIX": "MX",
    }
    resolved = brand
    if brand == "MIX":
        low = name.lower()
        if "moto" in low:
            resolved = "MOTO"
        elif "asus" in low:
            resolved = "ASUS"
    pre = prefix_map.get(brand, prefix_map.get(resolved, "SKU"))[:3]
    suf = ""
    if quality == "CareOG":
        suf = "CARE"
    elif quality == "With Frame":
        suf = "WF"
    elif screen == "OLED":
        suf = "OLED"
    else:
        suf = "LCD"
    base = (pre + clean[:8] + suf)[:15]
    sku = base
    i = 2
    while sku.upper() in existing:
        sku = f"{base[:12]}{i}"
        i += 1
    existing.add(sku.upper())
    return sku


def index_category_report(ws) -> tuple[dict[str, int], dict[str, int]]:
    by_title: dict[str, int] = {}
    by_sku: dict[str, int] = {}
    for row in range(7, ws.max_row + 1):
        sku = ws.cell(row, COL_SKU).value
        if sku:
            by_sku[str(sku).strip().lower()] = row
        for col in (COL_TITLE, COL_ITEM_NAME):
            title = ws.cell(row, col).value
            if title:
                key = str(title).strip().lower()[:80]
                by_title.setdefault(key, row)
    return by_title, by_sku


def find_donor_row(ws, sample_title: str, brand: str, name: str, by_title: dict[str, int]) -> int | None:
    if sample_title:
        key = sample_title.strip().lower()[:80]
        if key in by_title:
            return by_title[key]
        for k, row in by_title.items():
            if key[:40] in k or k[:40] in key:
                return row
    # fallback: search by model tokens in item name
    tokens = [t for t in re.sub(r"[^a-z0-9+ ]", " ", name.lower()).split() if len(t) > 2]
    best_row, best_score = None, 0
    for row in range(7, min(ws.max_row + 1, 1300)):
        title = str(ws.cell(row, COL_ITEM_NAME).value or ws.cell(row, COL_TITLE).value or "").lower()
        if not title:
            continue
        score = sum(1 for t in tokens if t in title)
        if score > best_score:
            best_score, best_row = score, row
    return best_row if best_score >= 2 else None


def clone_row(ws, src_row: int, dst_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        dst.value = src.value
        if src.has_style:
            dst._style = copy(src._style)


def audit_row(title: str, bullets: list[str]) -> list[str]:
    flags = []
    if len(title) > 200:
        flags.append("title>200")
    for i, b in enumerate(bullets, 1):
        if len(b) > 255:
            flags.append(f"bullet{i}>255")
    banned = ["best", "number1", "100% original", "whatsapp", "gst invoice"]
    text = (title + " " + " ".join(bullets)).lower()
    for term in banned:
        if term in text:
            flags.append(f"banned:{term}")
    return flags


def main() -> None:
    print("=" * 70)
    print("BULK CREATE — 25 missing quality listings")
    print("=" * 70)

    markups = compute_markups()
    print("Markup multipliers:", markups)

    sunsky_idx = load_sunsky_index()
    existing_skus = load_existing_skus()

    with MISSING_CSV.open(newline="", encoding="utf-8") as f:
        targets = list(csv.DictReader(f))
    print(f"Targets: {len(targets)}")

    cat_wb = openpyxl.load_workbook(CAT_REPORT)
    cat_ws = cat_wb.active
    max_col = cat_ws.max_column

    out_wb = openpyxl.load_workbook(CAT_REPORT)
    out_ws = out_wb.active
    # clear old data rows, keep header rows 1-6
    if out_ws.max_row > 6:
        out_ws.delete_rows(7, out_ws.max_row - 6)

    by_title, _ = index_category_report(cat_ws)
    review_rows = []
    out_row = 7

    for item in targets:
        brand = item["brand"]
        name = item["name"]
        qty = int(item.get("qty") or 1)
        screen = item["stock_screen"]
        quality = item["stock_quality"]
        sample_title = (item.get("sample_title") or "").strip()

        sunsky = sunsky_idx.get(name, {})

        donor_row = find_donor_row(cat_ws, sample_title, brand, name, by_title)
        flags: list[str] = []
        if not donor_row:
            flags.append("no_donor_row")
            donor_row = 7

        clone_row(cat_ws, donor_row, out_row, max_col)
        donor_sku = str(cat_ws.cell(donor_row, COL_SKU).value or "")

        title = build_title(brand, name, screen, quality)
        model = clean_model_name(brand_prefix(brand, name))
        bullets = make_bullets(title, screen, quality, model)
        keywords = make_keywords(screen, quality, model)
        new_sku = make_sku(brand, name, screen, quality, existing_skus)

        sell_price, sunsky_cost, price_flags = resolve_sell_price(
            name, screen, quality, sunsky, markups
        )
        flags.extend(price_flags)
        tk = tier_key(screen, quality)
        markup = markups.get(tk, DEFAULT_MARKUP[tk])
        mrp = calculate_mrp(sell_price) if sell_price else 0

        flags.extend(audit_row(title, bullets))

        out_ws.cell(out_row, COL_STATUS).value = None
        out_ws.cell(out_row, COL_SKU).value = new_sku
        out_ws.cell(out_row, COL_PRODUCT_TYPE).value = "PHONE_ACCESSORY"
        out_ws.cell(out_row, COL_LISTING_ACTION).value = "Create or Replace (Full Update)"
        out_ws.cell(out_row, COL_ITEM_NAME).value = title
        out_ws.cell(out_row, COL_TITLE).value = title
        out_ws.cell(out_row, COL_BRAND).value = "Generic"
        out_ws.cell(out_row, COL_DESCRIPTION).value = (
            "Compatible mobile phone display screen replacement. Tested for quality and performance "
            "before dispatch. Professional installation recommended for best results."
        )
        for i, b in enumerate(bullets):
            out_ws.cell(out_row, COL_BULLET1 + i).value = b
        out_ws.cell(out_row, COL_KEYWORD1).value = keywords
        out_ws.cell(out_row, COL_QTY).value = max(qty, 100)
        if sell_price:
            out_ws.cell(out_row, COL_PRICE).value = sell_price
            out_ws.cell(out_row, COL_MRP).value = mrp

        review_rows.append({
            "brand": brand,
            "stock_name": name,
            "new_sku": new_sku,
            "title": title,
            "sunsky_cost": sunsky_cost or "",
            "markup": markup,
            "sell_price": sell_price or "",
            "mrp": mrp or "",
            "qty": max(qty, 100),
            "donor_sku": donor_sku,
            "donor_title": str(cat_ws.cell(donor_row, COL_ITEM_NAME).value or "")[:80],
            "flags": " | ".join(flags),
        })

        print(f"  {new_sku}: {name[:40]} -> Rs.{sell_price:.0f}" if sell_price else f"  {new_sku}: {name[:40]} (NO PRICE)")
        out_row += 1

    out_wb.save(OUTPUT_XLSM)
    out_wb.close()
    cat_wb.close()

    with REVIEW_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(review_rows[0].keys()))
        w.writeheader()
        w.writerows(review_rows)

    flagged = [r for r in review_rows if r["flags"]]
    print(f"\nSaved: {OUTPUT_XLSM.name} ({len(review_rows)} rows)")
    print(f"Review: {REVIEW_CSV.name}")
    print(f"Flagged rows: {len(flagged)}")
    for r in flagged:
        print(f"  ! {r['stock_name']}: {r['flags']}")


if __name__ == "__main__":
    main()
