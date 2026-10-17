"""Fill row 7 upload-ready — official template + all Sir/sample fields. Excel COM only."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
import win32com.client as win32

BACKUP = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_backup.xlsm")
OUT = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm")
CAT = Path(
    r"C:\Projects\Amazon Systems Design\catalog-builder\amazon-reports"
    r"\Category+Listings+Report_05-12-2026.xlsm"
)
SUMMARY = Path(__file__).parent / "PHONE_ACCESSORY_upload_ready.txt"
DATA_ROW = 7
DONOR_ROW = 56
IMAGE_COLS = list(range(21, 31))


def browse_mobile_displays(template: Path) -> str:
    wb = openpyxl.load_workbook(template, read_only=True, data_only=True)
    vv = wb["Valid Values"]
    for c in range(3, vv.max_column + 1):
        v = vv.cell(10, c).value
        if v and "Mobile Displays" in str(v) and "11180541031" in str(v):
            wb.close()
            return str(v)
    wb.close()
    raise ValueError("Mobile Displays browse node not found")


def donor_text(col: int) -> str | None:
    wb = openpyxl.load_workbook(CAT, read_only=True, data_only=True)
    v = wb["Template"].cell(DONOR_ROW, col).value
    wb.close()
    return str(v) if v is not None else None


def make_bullets(title: str) -> list[str]:
    t = title.lower()
    tier = "gold" if "careog" in t else "silver" if "oled" in t else "bronze"
    fp = "fingerprint support" in t and "no fingerprint" not in t
    frame = "with frame" in t or "folder" in t
    m = re.search(r"compatible for\s+(.+?)(?:\s+\(|display)", title, re.I)
    model = m.group(1).strip() if m else "your phone"
    b1 = (
        f"Compatible with {model}"
        + (" — Fingerprint supported after fitting" if fp else " — Fingerprint not supported")
        + " — Verify model number from Settings before ordering"
    )
    tier_txt = {
        "gold": "Gold Tier CareOG — Tested OG combo with premium quality for professional repair.",
        "silver": "Silver Tier OLED — Premium display with vivid colors and better contrast.",
        "bronze": "Bronze Tier LCD — Budget friendly good quality display. Reliable performance.",
    }
    b3 = (
        "Quality Tested — Checked for dead pixels, color calibration, touch response."
        + (" With frame for easy installation." if frame else "")
    )
    return [
        b1[:255],
        tier_txt[tier][:255],
        b3[:255],
        "Professional installation recommended — Self-install without skills may damage device.",
        "7-Day replacement warranty — QC tested before dispatch. Film removal = acceptance.",
    ]


def build_cells() -> dict[int, object]:
    sku = "SAMPLE-S21FE-OLED"
    title = (
        "Compatible for Samsung Galaxy S21 FE (Fingerprint Support) "
        "OLED Display+Touch Screen Combo Folder"
    )
    bullets = make_bullets(title)
    browse = browse_mobile_displays(BACKUP)
    packer = donor_text(121)
    warranty = donor_text(470)

    cells: dict[int, object] = {
        1: sku,
        2: "PHONE_ACCESSORY",
        3: "Create or Replace (Full Update)",
        7: title,
        9: "Generic",
        10: "GTIN Exempt",
        12: browse,
        17: sku,
        18: "Generic",
        31: (
            "Compatible mobile phone OLED display and touch screen combo replacement. "
            "Tested for quality before dispatch. Professional installation recommended."
        ),
        37: "galaxy s21 fe oled display touch screen combo folder replacement",
        39: 1,
        40: "Mobile Screen Replacement Part",
        43: sku,
        45: "Samsung Galaxy S21 FE",
        49: 1,
        50: "Piece",
        52: "Mobile Screen",
        57: "HSN Code",
        58: "851770",
        63: packer,
        68: packer,
        74: 25,
        75: "Grams",
        77: "New",
        79: "A_GEN_STANDARD",
        81: 1,
        101: "AMAZON_IN",
        102: 1,
        103: 2,
        105: "Disabled",
        106: 2586,
        107: 3700,
        134: 26,
        135: "Centimetres",
        136: 12,
        137: "Centimetres",
        138: 2,
        139: "Centimetres",
        140: 26,
        141: "Centimetres",
        142: 12,
        143: "Centimetres",
        144: 2,
        145: "Centimetres",
        146: 220,
        147: "Grams",
        149: "China",
        150: warranty,
        153: "No",
        176: "Not Applicable",
    }
    for i, b in enumerate(bullets):
        cells[32 + i] = b
    return cells


def main() -> None:
    shutil.copy2(BACKUP, OUT)
    cells = build_cells()

    excel = win32.DispatchEx("Excel.Application")
    try:
        excel.DisplayAlerts = False
    except AttributeError:
        pass
    try:
        wb = excel.Workbooks.Open(str(OUT.resolve()))
        ws = wb.Worksheets("Template")
        # B before C/L — dropdown dependency
        ws.Cells(DATA_ROW, 2).Value = cells[2]
        ws.Cells(DATA_ROW, 3).Value = cells[3]
        ws.Cells(DATA_ROW, 12).Value = cells[12]
        for c in IMAGE_COLS:
            ws.Cells(DATA_ROW, c).ClearContents()
        skip = {2, 3, 12, *IMAGE_COLS}
        for col, val in sorted(cells.items()):
            if col in skip or val is None or str(val).strip() == "":
                continue
            ws.Cells(DATA_ROW, col).Value = val
        wb.Save()
        wb.Close(SaveChanges=True)
    finally:
        excel.Quit()

    lines = [f"Upload-ready: {OUT}", f"Size: {OUT.stat().st_size:,}", f"Row: {DATA_ROW}", ""]
    wb2 = openpyxl.load_workbook(OUT, read_only=True, data_only=True)
    ws2 = wb2["Template"]
    lines.append(f"r6 C (untouched): {ws2.cell(6, 3).value!r}")
    lines.append(f"Images 21-30 blank: {all(not ws2.cell(7, c).value for c in IMAGE_COLS)}")
    lines.append("\nFilled:")
    for c in range(1, ws2.max_column + 1):
        v = ws2.cell(DATA_ROW, c).value
        if v is not None and str(v).strip():
            lines.append(f"  col{c:3d} {ws2.cell(4, c).value}: {str(v)[:80]}")
    wb2.close()
    SUMMARY.write_text("\n".join(lines), encoding="utf-8")
    print(SUMMARY)


if __name__ == "__main__":
    main()
