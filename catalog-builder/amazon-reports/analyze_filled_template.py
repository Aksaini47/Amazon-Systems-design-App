"""Analyse filled PHONE_ACCESSORY template — required vs filled columns."""
from __future__ import annotations

from pathlib import Path

import openpyxl

TEMPLATE = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm")
OUT = Path(__file__).parent / "PHONE_ACCESSORY_fill_analysis.txt"
DATA_ROW = 7

# Image columns (new 237-col template)
IMAGE_COLS = list(range(21, 30)) + [30]  # 21-29 other/main, 30 swatch

# Mandatory for CREATE (from Data Definitions + Amazon PHONE_ACCESSORY rules)
MANDATORY_CREATE = {
    1: "SKU",
    2: "Product Type",
    3: "Listing Action",
    7: "Item Name",
    9: "Brand Name",
    12: "Recommended Browse Nodes",
    101: "Fulfillment Channel Code (IN)",
    102: "Quantity (IN)",
}

# Strongly recommended for CREATE (will fail without in practice)
RECOMMENDED_CREATE = {
    21: "Main Image URL",
    31: "Product Description",
    32: "Bullet Point 1",
    106: "Your Price INR",
    107: "Maximum Retail Price",
    77: "Item Condition",
    58: "External Product Information (HSN)",
    57: "External Product Information Entity",
    149: "Country of Origin",
    153: "Are batteries required?",
}


def main() -> None:
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)
    ws = wb["Template"]
    dd = wb["Data Definitions"]
    vv = wb["Valid Values"]

    # Listing actions from Valid Values row 3
    listing_actions = [
        vv.cell(3, c).value
        for c in range(3, 10)
        if vv.cell(3, c).value
        and str(vv.cell(3, c).value) in (
            "Create or Replace (Full Update)",
            "Edit (Partial Update)",
            "Delete",
        )
    ]

    # Build required map from Data Definitions
    dd_required: dict[int, str] = {}
    field_to_col: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        label = ws.cell(4, c).value
        field = ws.cell(5, c).value
        if label and field:
            field_to_col[str(field)] = c

    for r in range(4, dd.max_row + 1):
        field = dd.cell(r, 2).value
        label = dd.cell(r, 3).value
        req = dd.cell(r, 6).value
        if field and req and str(req).strip().lower() == "required":
            for c in range(1, ws.max_column + 1):
                if ws.cell(5, c).value == field:
                    dd_required[c] = str(label or field)
                    break

    filled = []
    blank = []
    for c in range(1, ws.max_column + 1):
        label = ws.cell(4, c).value or f"col{c}"
        v = ws.cell(DATA_ROW, c).value
        if v is not None and str(v).strip():
            filled.append((c, str(label)[:50], str(v)[:90]))
        else:
            blank.append((c, str(label)[:50]))

    lines = [
        f"File: {TEMPLATE}",
        f"Data row: {DATA_ROW}",
        f"Total cols: {ws.max_column}",
        f"Filled: {len(filled)} | Blank: {len(blank)}",
        "",
        f"Valid Listing Actions: {listing_actions}",
        f"Row7 Col C (Listing Action): {ws.cell(DATA_ROW, 3).value!r}",
        f"Row6 Col C (template default): {ws.cell(6, 3).value!r}",
        f"Row4 Col C (header): {ws.cell(4, 3).value!r}",
        "",
        "=== MANDATORY (CREATE) — check ===",
    ]
    for col, name in sorted(MANDATORY_CREATE.items()):
        v = ws.cell(DATA_ROW, col).value
        status = "OK" if v is not None and str(v).strip() else "MISSING"
        lines.append(f"  [{status}] col{col:3d} {name}: {str(v)[:80] if v else '-'}")

    lines.append("\n=== RECOMMENDED (CREATE) — check ===")
    for col, name in sorted(RECOMMENDED_CREATE.items()):
        v = ws.cell(DATA_ROW, col).value
        status = "filled" if v is not None and str(v).strip() else "blank"
        lines.append(f"  [{status}] col{col:3d} {name}")

    lines.append(f"\n=== DATA DEFINITIONS 'Required' ({len(dd_required)} fields) ===")
    for col in sorted(dd_required):
        v = ws.cell(DATA_ROW, col).value
        st = "OK" if v is not None and str(v).strip() else "MISSING"
        lines.append(f"  [{st}] col{col:3d} {dd_required[col]}")

    lines.append("\n=== FILLED BY SIR (all non-empty row 7) ===")
    for c, lbl, val in filled:
        tag = ""
        if c in MANDATORY_CREATE:
            tag = " [MANDATORY]"
        elif c in RECOMMENDED_CREATE:
            tag = " [recommended]"
        elif c in IMAGE_COLS:
            tag = " [IMAGE]"
        lines.append(f"  col{c:3d} {lbl}{tag}: {val}")

    lines.append(f"\n=== BLANK — not mandatory (sample, first 40) ===")
    non_mandatory_blank = [
        (c, lbl)
        for c, lbl in blank
        if c not in MANDATORY_CREATE and c not in dd_required
    ]
    lines.append(f"  Count: {len(non_mandatory_blank)} blank optional cols")
    for c, lbl in non_mandatory_blank[:40]:
        lines.append(f"  col{c:3d} {lbl}")
    if len(non_mandatory_blank) > 40:
        lines.append(f"  ... +{len(non_mandatory_blank)-40} more")

    wb.close()
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
