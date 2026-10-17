import openpyxl
from pathlib import Path

def rows_1_6(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Template"]
    out = []
    for r in range(1, 7):
        row = {}
        for c in range(1, 20):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                row[c] = str(v)[:60]
        if row:
            out.append((r, row))
    vv = wb["Valid Values"]
    vv3 = [vv.cell(3, c).value for c in range(2, 7)]
    dl = [wb["Dropdown Lists"].cell(r, 3).value for r in range(4, 7)]
    wb.close()
    return out, vv3, dl

for label, p in [("backup", Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_backup.xlsm")),
                 ("current", Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm"))]:
    r6, vv3, dl = rows_1_6(p)
    print("===", label)
    print("VV row3:", vv3)
    print("DL C4-6:", dl)
    for r, row in r6:
        if 3 in row or 12 in row or r == 6:
            print(f" row{r} C={row.get(3)} L={row.get(12)}")
