"""Thorough analysis of PHONE_ACCESSORY template sheets."""
from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl

BACKUP = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_backup.xlsm")
CURRENT = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm")
OUT = Path(__file__).parent / "_template_full_analysis.txt"


def analyze(path: Path) -> list[str]:
    lines = [f"\n{'='*60}", f"FILE: {path.name} ({path.stat().st_size:,} bytes)"]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    for sn in wb.sheetnames:
        ws = wb[sn]
        lines.append(f"\n--- Sheet: {sn} (rows={ws.max_row}, cols={ws.max_column}) ---")
        if ws.max_row >= 4:
            r4 = [ws.cell(4, c).value for c in range(1, min(20, ws.max_column + 1))]
            lines.append(f"  row4[:19]: {r4}")
        if sn == "Valid Values" and ws.max_row > 1:
            for r in range(1, min(30, ws.max_row + 1)):
                row = [ws.cell(r, c).value for c in range(1, min(8, ws.max_column + 1))]
                if any(row):
                    lines.append(f"  r{r}: {row}")
        if sn == "Dropdown Lists" and ws.max_row > 1:
            for r in range(1, min(40, ws.max_row + 1)):
                row = [ws.cell(r, c).value for c in range(1, min(6, ws.max_column + 1))]
                if any(row):
                    lines.append(f"  r{r}: {row}")
        if sn == "Browse Data" and ws.max_row > 1:
            for r in range(1, min(25, ws.max_row + 1)):
                row = [ws.cell(r, c).value for c in range(1, min(4, ws.max_column + 1))]
                if any(row):
                    lines.append(f"  r{r}: {row}")
        if sn == "Data Definitions" and ws.max_row > 1:
            for r in range(1, min(25, ws.max_row + 1)):
                row = [ws.cell(r, c).value for c in range(1, min(6, ws.max_column + 1))]
                if any(row):
                    lines.append(f"  r{r}: {row}")
        if sn == "Template":
            lines.append("  Template row5 cols 1-16:")
            for c in range(1, 17):
                lines.append(f"    col{c}: r4={ws.cell(4,c).value!r} r5={str(ws.cell(5,c).value or '')[:70]!r}")
            lines.append("  Template row7 cols 1-16:")
            for c in range(1, 17):
                v = ws.cell(7, c).value
                if v is not None:
                    lines.append(f"    col{c}: {str(v)[:80]!r}")

    wb.close()

    # data validations in sheet5.xml
    with zipfile.ZipFile(path) as z:
        if "xl/worksheets/sheet5.xml" in z.namelist():
            xml = z.read("xl/worksheets/sheet5.xml").decode("utf-8", errors="replace")
            if "dataValidation" in xml:
                import re
                dvs = re.findall(r'<dataValidation[^>]*sqref="([^"]*)"[^>]*(?:/>|>.*?</dataValidation>)', xml, re.S)
                lines.append(f"\n  sheet5 dataValidations count: {len(dvs)}")
                for i, block in enumerate(re.findall(r"<dataValidation[^>]*>.*?</dataValidation>|<dataValidation[^>]*/>", xml, re.S)[:8]):
                    sq = re.search(r'sqref="([^"]*)"', block)
                    f1 = re.search(r'formula1="([^"]*)"', block)
                    t = re.search(r'type="([^"]*)"', block)
                    lines.append(f"    dv{i}: type={t.group(1) if t else '?'} sqref={sq.group(1) if sq else '?'} formula1={f1.group(1)[:60] if f1 else '?'}")
    return lines


def main() -> None:
    parts = []
    for p in (BACKUP, CURRENT):
        if p.exists():
            parts.extend(analyze(p))
    OUT.write_text("\n".join(parts), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
