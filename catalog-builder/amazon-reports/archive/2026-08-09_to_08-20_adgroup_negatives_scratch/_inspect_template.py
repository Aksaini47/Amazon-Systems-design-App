"""Inspect new PHONE_ACCESSORY template vs category report."""
import openpyxl
from pathlib import Path

TEMPLATE = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm")
CAT = Path(r"C:\Projects\Amazon Systems Design\catalog-builder\amazon-reports\Category+Listings+Report_05-12-2026.xlsm")
OUT = Path(__file__).parent / "_template_inspect.txt"


def col_map(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Template"]
    m = {}
    for c in range(1, ws.max_column + 1):
        label = ws.cell(4, c).value
        field = ws.cell(5, c).value
        if label:
            m[c] = {"label": str(label), "field": str(field or "")[:80]}
    wb.close()
    return m


keys = [
    "sku", "product type", "listing action", "item name", "brand", "image",
    "description", "bullet", "keyword", "browse", "fulfillment", "quantity",
    "price", "maximum retail", "compatible", "condition", "hsn", "batter",
    "product id", "manufacturer", "model", "color", "material", "part",
    "warranty", "country", "item type", "supplier",
]

lines = []
tm = col_map(TEMPLATE)
lines.append("=== NEW TEMPLATE (237 cols) ===")
for c, d in sorted(tm.items()):
    l = d["label"].lower()
    if any(k in l for k in keys):
        lines.append(f"{c:3d} | {d['label'][:60]}")

wb = openpyxl.load_workbook(CAT, read_only=True, data_only=True)
ws = wb["Template"]
cat_labels = {
    str(ws.cell(4, c).value).strip(): c
    for c in range(1, ws.max_column + 1)
    if ws.cell(4, c).value
}

donor = None
for r in range(7, ws.max_row + 1):
    pt = str(ws.cell(r, 4).value or "")
    sku = ws.cell(r, 3).value
    if pt == "PHONE_ACCESSORY" and sku:
        donor = r
        break

lines.append(f"\ndonor row={donor} sku={ws.cell(donor, 3).value if donor else None}")

search_labels = [
    "SKU", "Product Type", "Listing Action", "Item Name", "Brand Name",
    "Main Image URL", "Product Description", "Bullet Point",
    "Generic Keyword", "Compatible Phone Models",
    "Fulfillment Channel Code (IN)", "Quantity (IN)",
    "Your Price INR", "Maximum Retail Price",
    "Recommended Browse Nodes", "Are batteries required?",
    "Condition Type", "Manufacturer Part Number", "Model Number",
    "Color", "Material Type", "Item Type Name", "Country of Origin",
]
if donor:
    lines.append("\n=== DONOR VALUES (CAT report) ===")
    for lbl in search_labels:
        matches = [(c, l) for l, c in cat_labels.items() if lbl.lower() in l.lower()]
        if matches:
            c = matches[0][0]
            v = ws.cell(donor, c).value
            if v is not None and str(v).strip():
                lines.append(f"col{c:3d} {matches[0][1][:50]} = {str(v)[:90]}")

# Also find PHONE_ACCESSORY in category - maybe different product type name
pt_counts = {}
for r in range(7, min(200, ws.max_row + 1)):
    pt = str(ws.cell(r, 4).value or "")
    pt_counts[pt] = pt_counts.get(pt, 0) + 1
lines.append("\nproduct types in first 200 rows: " + str(pt_counts))
wb.close()

OUT.write_text("\n".join(lines), encoding="utf-8")
print(OUT)
