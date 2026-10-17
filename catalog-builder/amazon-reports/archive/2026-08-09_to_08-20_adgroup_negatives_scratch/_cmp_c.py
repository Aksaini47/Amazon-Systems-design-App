import openpyxl
import zipfile
from pathlib import Path

for label, p in [
    ("backup", Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_backup.xlsm")),
    ("current", Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4).xlsm")),
]:
    print("===", label, p.stat().st_size)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb["Template"]
    for r in range(4, 8):
        print(f"  r{r} B={ws.cell(r,2).value!r} C={ws.cell(r,3).value!r}")
    dl = wb["Dropdown Lists"]
    print("  DL C3-C8:", [dl.cell(r, 3).value for r in range(3, 9)])
    vv = wb["Valid Values"]
    print("  VV r3 c2-6:", [vv.cell(3, c).value for c in range(2, 7)])
    wb.close()
    with zipfile.ZipFile(p) as z:
        x = z.read("xl/worksheets/sheet5.xml").decode()
        i = x.find('sqref="C7:C1048576"')
        if i > 0:
            j = x.find("dataValidation", i - 200)
            print("  C7 dv:", x[j : i + 120])
