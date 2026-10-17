"""Fill one sample CREATE row in PHONE_ACCESSORY (4).xlsm — mandatory + relevant fields only."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl

SOURCE = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_backup.xlsm")
# Write to new file — openpyxl must use keep_vba=True or .xlsm macros break
OUT_FILE = Path(r"C:\Users\DELL\Downloads\PHONE_ACCESSORY (4)_sample.xlsm")
PROJECT_COPY = Path(__file__).parent / "PHONE_ACCESSORY_sample_row.xlsm"
CAT = Path(
    r"C:\Projects\Amazon Systems Design\catalog-builder\amazon-reports"
    r"\Category+Listings+Report_05-12-2026.xlsm"
)
DATA_ROW = 7
DONOR_ROW = 56  # Poco C31 CareOG — PHONE_ACCESSORY reference

# New template (237 cols) — row 4 headers
T = {
    "sku": 1,
    "product_type": 2,
    "listing_action": 3,
    "item_name": 7,
    "brand": 9,
    "browse1": 12,
    "model_number": 17,
    "manufacturer": 18,
    "main_image": 21,
    "description": 31,
    "bullet1": 32,
    "keyword": 37,
    "num_items": 39,
    "item_type": 40,
    "part_number": 43,
    "compat1": 45,
    "hsn_entity": 57,
    "hsn": 58,
    "packer": 68,
    "condition": 77,
    "tax_code": 79,
    "max_order_qty": 81,
    "fulfillment": 101,
    "quantity": 102,
    "price": 106,
    "mrp": 107,
    "pkg_len": 140,
    "pkg_len_u": 141,
    "pkg_w": 142,
    "pkg_w_u": 143,
    "pkg_h": 144,
    "pkg_h_u": 145,
    "pkg_weight": 146,
    "pkg_weight_u": 147,
    "country": 149,
    "warranty": 150,
    "batteries_required": 153,
    "dangerous_goods": 176,
}

# Category report (591 cols) — donor field locations
C = {
    "browse1": 13,
    "manufacturer": 20,
    "main_image": 29,
    "num_items": 61,
    "item_type": 63,
    "hsn_entity": 114,
    "hsn": 115,
    "packer": 121,
    "condition": 395,
    "tax_code": 397,
    "max_order_qty": 399,
    "fulfillment": 419,
    "pkg_len": 452,
    "pkg_len_u": 453,
    "pkg_w": 454,
    "pkg_w_u": 455,
    "pkg_h": 456,
    "pkg_h_u": 457,
    "pkg_weight": 458,
    "pkg_weight_u": 459,
    "country": 469,
    "warranty": 470,
    "batteries_required": 475,
    "dangerous_goods": 498,
}


def donor_val(cat_ws, col: int):
    return cat_ws.cell(DONOR_ROW, col).value


def make_bullets(title: str) -> list[str]:
    t = title.lower()
    tier = "gold" if "careog" in t else "silver" if "oled" in t else "bronze"
    fp = "fingerprint support" in t and "no fingerprint" not in t
    frame = "with frame" in t or "folder" in t
    m = re.search(r"compatible for\s+(.+?)(?:\s+\(|display)", title, re.I)
    model = m.group(1).strip() if m else "your phone"
    b1 = (
        f"Compatible with {model}"
        + (" — Fingerprint supported after fitting" if fp else " — Fingerprint not supported")
        + " — Verify model number from Settings before ordering"
    )
    tier_txt = {
        "gold": "Gold Tier CareOG — Tested OG combo with premium quality for professional repair.",
        "silver": "Silver Tier OLED — Premium display with vivid colors and better contrast.",
        "bronze": "Bronze Tier LCD — Budget friendly good quality display. Reliable performance.",
    }
    b3 = (
        "Quality Tested — Checked for dead pixels, color calibration, touch response."
        + (" With frame for easy installation." if frame else "")
    )
    return [
        b1[:255],
        tier_txt[tier][:255],
        b3[:255],
        "Professional installation recommended — Self-install without skills may damage device.",
        "7-Day replacement warranty — QC tested before dispatch. Film removal = acceptance.",
    ]


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Backup missing: {SOURCE}")

    cat_wb = openpyxl.load_workbook(CAT, read_only=True, data_only=True)
    cat_ws = cat_wb["Template"]

    new_sku = "SAMPLE-S21FE-OLED"
    new_title = (
        "Compatible for Samsung Galaxy S21 FE (Fingerprint Support) "
        "OLED Display+Touch Screen Combo Folder"
    )
    sell_price = 2586
    mrp = 3700
    bullets = make_bullets(new_title)

    wb = openpyxl.load_workbook(SOURCE, keep_vba=True)
    ws = wb["Template"]

    def setf(key: str, value) -> None:
        if value is not None and str(value).strip() != "":
            ws.cell(DATA_ROW, T[key]).value = value

    # Mandatory
    setf("sku", new_sku)
    setf("product_type", "PHONE_ACCESSORY")
    setf("listing_action", "Create or Replace (Full Update)")
    setf("item_name", new_title)
    setf("brand", "Generic")
    setf("fulfillment", donor_val(cat_ws, C["fulfillment"]) or "Fulfillment by Merchant (Default)")
    setf("quantity", 100)

    # Relevant create fields (from donor format)
    setf("browse1", donor_val(cat_ws, C["browse1"]))
    setf("model_number", new_sku)
    setf("manufacturer", donor_val(cat_ws, C["manufacturer"]) or "Generic")
    setf("main_image", donor_val(cat_ws, C["main_image"]))
    setf(
        "description",
        "Compatible mobile phone OLED display and touch screen combo replacement. "
        "Tested for quality before dispatch. Professional installation recommended.",
    )
    for i, b in enumerate(bullets):
        ws.cell(DATA_ROW, T["bullet1"] + i).value = b
    setf("keyword", "galaxy s21 fe oled display touch screen combo folder replacement")
    setf("num_items", donor_val(cat_ws, C["num_items"]) or 1)
    setf("item_type", donor_val(cat_ws, C["item_type"]) or "Mobile Screen Replacement Part")
    setf("part_number", new_sku)
    setf("compat1", "Samsung Galaxy S21 FE")
    setf("hsn_entity", donor_val(cat_ws, C["hsn_entity"]) or "HSN Code")
    setf("hsn", donor_val(cat_ws, C["hsn"]) or "851770")
    setf("packer", donor_val(cat_ws, C["packer"]))
    setf("condition", donor_val(cat_ws, C["condition"]) or "New")
    setf("tax_code", donor_val(cat_ws, C["tax_code"]) or "A_ELECTRONICS_GEN")
    setf("max_order_qty", donor_val(cat_ws, C["max_order_qty"]) or 1)
    setf("price", sell_price)
    setf("mrp", mrp)
    for k in ("pkg_len", "pkg_len_u", "pkg_w", "pkg_w_u", "pkg_h", "pkg_h_u", "pkg_weight", "pkg_weight_u"):
        setf(k, donor_val(cat_ws, C[k]))
    setf("country", donor_val(cat_ws, C["country"]) or "China")
    setf("warranty", donor_val(cat_ws, C["warranty"]))
    setf("batteries_required", donor_val(cat_ws, C["batteries_required"]) or "No")
    setf("dangerous_goods", donor_val(cat_ws, C["dangerous_goods"]) or "Not Applicable")

    cat_wb.close()
    wb.save(OUT_FILE)
    wb.close()
    shutil.copy2(OUT_FILE, PROJECT_COPY)

    filled = []
    wb2 = openpyxl.load_workbook(OUT_FILE, read_only=True, data_only=True)
    ws2 = wb2["Template"]
    for c in range(1, ws2.max_column + 1):
        v = ws2.cell(DATA_ROW, c).value
        if v is not None and str(v).strip():
            filled.append((c, ws2.cell(4, c).value, str(v)[:70]))
    wb2.close()

    summary = Path(__file__).parent / "PHONE_ACCESSORY_sample_row_summary.txt"
    lines = [
        f"Output: {OUT_FILE}",
        f"Source (intact): {SOURCE}",
        f"Row: {DATA_ROW}",
        f"SKU: {new_sku}",
        f"Action: Create or Replace (Full Update)",
        f"Title: {new_title}",
        f"Price: {sell_price} | MRP: {mrp} | Qty: 100",
        f"Donor reference: CAT row {DONOR_ROW} (Poco C31 CareOG)",
        "",
        "Filled columns:",
    ]
    lines += [f"  col{c:3d} {lbl}: {val}" for c, lbl, val in filled]
    summary.write_text("\n".join(lines), encoding="utf-8")
    print(summary)
    print(f"Filled {len(filled)} columns in row {DATA_ROW}")


if __name__ == "__main__":
    main()
